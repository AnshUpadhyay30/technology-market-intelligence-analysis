import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.analysis import (
    build_priority_analysis,
    build_competitor_analysis,
    build_statistical_summary,
    build_research_source_analysis,
)


REPORTS_DIR = "reports"
CHARTS_DIR = "charts"
OUTPUT_FILE = f"{REPORTS_DIR}/technology_market_intelligence_report_v2.xlsx"


COLORS = {
    "dark_blue": "1F4E78",
    "blue": "2F75B5",
    "sky_blue": "00A2E8",
    "light_blue": "D9EAF7",
    "very_light_blue": "EAF3F8",
    "green": "00A65A",
    "dark_green": "0B6E0B",
    "light_green": "D9EAD3",
    "orange": "E69138",
    "light_orange": "FCE5CD",
    "red": "E31B23",
    "light_red": "F4CCCC",
    "yellow": "FFF2CC",
    "gray": "6C7A89",
    "light_gray": "D9E2F3",
    "white": "FFFFFF",
    "black": "000000",
}


def ensure_reports_dir():
    os.makedirs(REPORTS_DIR, exist_ok=True)


def create_border():
    return Border(
        left=Side(style="thin", color="B7C9D6"),
        right=Side(style="thin", color="B7C9D6"),
        top=Side(style="thin", color="B7C9D6"),
        bottom=Side(style="thin", color="B7C9D6"),
    )


def set_page_style(ws):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"


def style_title(ws, title, subtitle, start_col=1, end_col=8):
    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
    ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=end_col)

    title_cell = ws.cell(row=1, column=start_col)
    subtitle_cell = ws.cell(row=2, column=start_col)

    title_cell.value = title
    subtitle_cell.value = subtitle

    title_cell.fill = PatternFill("solid", fgColor=COLORS["dark_blue"])
    title_cell.font = Font(color=COLORS["white"], bold=True, size=16)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    subtitle_cell.fill = PatternFill("solid", fgColor=COLORS["blue"])
    subtitle_cell.font = Font(color=COLORS["white"], bold=True, italic=True, size=10)
    subtitle_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 22


def style_headers(ws, header_row, start_col, end_col, fill_color=None):
    fill_color = fill_color or COLORS["dark_blue"]

    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.font = Font(color=COLORS["white"], bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = create_border()


def style_body(ws, min_row, max_row, min_col, max_col, fill_alternate=True):
    for row in range(min_row, max_row + 1):
        fill = COLORS["very_light_blue"] if row % 2 == 0 else COLORS["white"]

        for col in range(min_col, max_col + 1):
            cell = ws.cell(row=row, column=col)

            if fill_alternate:
                cell.fill = PatternFill("solid", fgColor=fill)

            cell.border = create_border()
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def auto_fit_columns(ws, min_width=10, max_width=32):
    for column_cells in ws.columns:
        col_letter = get_column_letter(column_cells[0].column)
        max_length = 0

        for cell in column_cells:
            if cell.value is not None:
                text = str(cell.value)
                if len(text) > max_length:
                    max_length = len(text)

        ws.column_dimensions[col_letter].width = min(max(max_length + 2, min_width), max_width)


def add_table_style(ws, table_name, start_row, start_col, end_row, end_col):
    if end_row <= start_row:
        return

    table_ref = (
        f"{get_column_letter(start_col)}{start_row}:"
        f"{get_column_letter(end_col)}{end_row}"
    )

    clean_name = table_name.replace(" ", "").replace("&", "").replace("-", "")[:25]

    try:
        table = Table(displayName=clean_name, ref=table_ref)
        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table.tableStyleInfo = style
        ws.add_table(table)
    except Exception:
        pass


def write_dataframe(
    ws,
    df,
    start_row,
    start_col,
    table_name=None,
    header_color=None,
    max_rows=None,
):
    if df is None or df.empty:
        ws.cell(start_row, start_col).value = "No data available"
        return start_row + 2

    if max_rows:
        df = df.head(max_rows)

    for col_idx, column_name in enumerate(df.columns, start=start_col):
        ws.cell(row=start_row, column=col_idx).value = column_name

    style_headers(
        ws,
        header_row=start_row,
        start_col=start_col,
        end_col=start_col + len(df.columns) - 1,
        fill_color=header_color or COLORS["dark_blue"],
    )

    for row_idx, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for col_idx, value in enumerate(row, start=start_col):
            ws.cell(row=row_idx, column=col_idx).value = value

    end_row = start_row + len(df)
    end_col = start_col + len(df.columns) - 1

    style_body(ws, start_row + 1, end_row, start_col, end_col)

    if table_name:
        add_table_style(ws, table_name, start_row, start_col, end_row, end_col)

    auto_fit_columns(ws)

    return end_row + 3


def apply_priority_colors(ws):
    priority_colors = {
        "Critical": COLORS["red"],
        "High": COLORS["orange"],
        "Medium": COLORS["green"],
        "Low": COLORS["gray"],
    }

    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row=row, column=col).value

            if value in priority_colors:
                cell = ws.cell(row=row, column=col)
                cell.fill = PatternFill("solid", fgColor=priority_colors[value])
                cell.font = Font(color=COLORS["white"], bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")


def format_date_columns(ws):
    for row in ws.iter_rows():
        for cell in row:
            if "date" in str(ws.cell(row=3, column=cell.column).value).lower():
                if cell.row > 3:
                    cell.number_format = "yyyy-mm-dd"


def create_executive_dashboard(wb, analysis_df, category_summary, top10_df, statistical_summary_df):
    ws = wb.create_sheet("📊 Executive Dashboard", 0)

    style_title(
        ws,
        "Technology Market Intelligence — Executive Dashboard",
        "R&D Priority & Strategic Technology Evaluation",
        start_col=1,
        end_col=8,
    )

    stats = statistical_summary_df.iloc[0].to_dict()

    kpis = [
        ("Total Technologies", stats.get("total_technologies", 0), COLORS["dark_blue"]),
        ("Avg Priority Score", stats.get("average_priority_score", 0), COLORS["blue"]),
        ("Avg Risk Score", stats.get("average_risk_score", 0), COLORS["red"]),
        ("Critical Priority", stats.get("critical_technologies", 0), COLORS["red"]),
        ("High Priority", stats.get("high_priority_technologies", 0), COLORS["orange"]),
        ("Medium Priority", stats.get("medium_priority_technologies", 0), COLORS["green"]),
        ("Low Priority", stats.get("low_priority_technologies", 0), COLORS["gray"]),
        ("Adoption Correlation", stats.get("adoption_priority_correlation", 0), COLORS["sky_blue"]),
    ]

    start_row = 4

    for index, (label, value, color) in enumerate(kpis, start=1):
        label_cell = ws.cell(row=start_row, column=index)
        label_cell.value = label
        label_cell.fill = PatternFill("solid", fgColor=color)
        label_cell.font = Font(color=COLORS["white"], bold=True, size=9)
        label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        label_cell.border = create_border()

        value_cell = ws.cell(row=start_row + 1, column=index)
        value_cell.value = value
        value_cell.fill = PatternFill("solid", fgColor=COLORS["light_blue"])
        value_cell.font = Font(color=color, bold=True, size=14)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.border = create_border()

        ws.column_dimensions[get_column_letter(index)].width = 18

    ws.row_dimensions[start_row].height = 22
    ws.row_dimensions[start_row + 1].height = 30

    ws.cell(row=7, column=1).value = "📌 Category Performance Summary"
    ws.cell(row=7, column=1).font = Font(bold=True, color=COLORS["dark_blue"], size=12)

    dashboard_category = category_summary[
        [
            "category",
            "total_technologies",
            "average_priority_score",
            "average_risk_score",
            "average_product_impact",
            "average_adoption_potential",
        ]
    ].copy()

    dashboard_category.columns = [
        "Category",
        "Technologies",
        "Avg Priority",
        "Avg Risk",
        "Avg Impact",
        "Avg Adoption",
    ]

    next_row = write_dataframe(
        ws,
        dashboard_category.head(12),
        start_row=8,
        start_col=1,
        table_name=None,
    )

    ws.cell(row=next_row + 1, column=1).value = "🏆 Top 10 Priority Technologies"
    ws.cell(row=next_row + 1, column=1).font = Font(
        bold=True, color=COLORS["dark_blue"], size=12
    )

    dashboard_top = top10_df[
        [
            "technology_name",
            "category",
            "final_priority_score",
            "priority_level",
            "recommendation",
        ]
    ].copy()

    dashboard_top.columns = [
        "Technology Name",
        "Category",
        "Priority Score",
        "Level",
        "Recommendation",
    ]

    write_dataframe(
        ws,
        dashboard_top,
        start_row=next_row + 2,
        start_col=1,
        table_name=None,
        header_color=COLORS["sky_blue"],
    )

    ws.freeze_panes = "A8"
    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 75

    apply_priority_colors(ws)


def create_priority_report(wb, priority_report):
    ws = wb.create_sheet("🎯 Priority Report")

    style_title(
        ws,
        "Technology Priority Report",
        "Ranked by Final Priority Score | Critical → High → Medium → Low",
        start_col=1,
        end_col=15,
    )

    df = priority_report.copy()
    df.columns = [
        "Technology ID",
        "Technology Name",
        "Category",
        "Expected Release Date",
        "Product Impact Score",
        "Adoption Potential Score",
        "Market Readiness Score",
        "R&D Investment Score",
        "Risk Score",
        "Competitor Activity Score",
        "Research Reliability Score",
        "Patent Reference Count",
        "Final Priority Score",
        "Priority Level",
        "Recommendation",
    ]

    write_dataframe(ws, df, start_row=3, start_col=1, table_name="PriorityReport")

    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False

    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["J"].width = 18
    ws.column_dimensions["K"].width = 18
    ws.column_dimensions["L"].width = 18
    ws.column_dimensions["O"].width = 80

    format_date_columns(ws)
    apply_priority_colors(ws)


def create_competitor_analysis(wb, competitor_summary_df):
    ws = wb.create_sheet("🏢 Competitor Analysis")

    style_title(
        ws,
        "Competitor Activity Analysis",
        "Market Adoption Tracking Across Industry Players",
        start_col=1,
        end_col=9,
    )

    df = competitor_summary_df.copy()
    df.columns = [
        "Competitor Name",
        "Total Tracked Technologies",
        "Average Market Activity Score",
        "Adopted Count",
        "Pilot Count",
        "Research Count",
        "Not Adopted Count",
        "High Investment Count",
        "Adoption Rate Percentage",
    ]

    write_dataframe(
        ws,
        df,
        start_row=3,
        start_col=1,
        table_name=None,
    )

    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False

    for row in range(4, ws.max_row + 1):
        adoption_cell = ws.cell(row=row, column=9)
        value = adoption_cell.value or 0

        if value >= 28:
            fill = COLORS["dark_green"]
        elif value >= 25:
            fill = "93C47D"
        else:
            fill = COLORS["light_red"]

        adoption_cell.fill = PatternFill("solid", fgColor=fill)
        adoption_cell.font = Font(bold=True, color=COLORS["black"])

    auto_fit_columns(ws, max_width=30)


def create_risk_analysis(wb, risk_report):
    ws = wb.create_sheet("⚠️ Risk Analysis")

    style_title(
        ws,
        "Technology Risk Analysis Report",
        "High Risk Technologies — Sorted by Risk Score",
        start_col=1,
        end_col=6,
    )

    df = risk_report.copy()
    df.columns = [
        "Technology Name",
        "Category",
        "Risk Score",
        "Final Priority Score",
        "Priority Level",
        "Recommendation",
    ]

    write_dataframe(
        ws,
        df.head(150),
        start_row=3,
        start_col=1,
        table_name=None,
        header_color=COLORS["red"],
    )

    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["F"].width = 85

    for row in range(4, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            if ws.cell(row=row, column=5).value not in ["Critical", "High", "Medium", "Low"]:
                ws.cell(row=row, column=col).fill = PatternFill(
                    "solid", fgColor=COLORS["light_red"]
                )

    apply_priority_colors(ws)


def create_patent_source_sheet(
    wb,
    source_type_summary,
    patent_summary,
    patent_office_summary,
    literature_summary,
    maturity_summary,
):
    ws = wb.create_sheet("📄 Patent & Source Intelligence")

    style_title(
        ws,
        "Patent Reference & Research Source Intelligence",
        "Source Reliability, Literature Coverage, Patent References, and Technology Maturity",
        start_col=1,
        end_col=8,
    )

    ws.cell(row=4, column=1).value = "📚 Source Type Summary"
    ws.cell(row=4, column=1).font = Font(bold=True, color=COLORS["dark_blue"], size=12)

    source_df = source_type_summary.copy()
    source_df.columns = ["Source Type", "Total Sources", "Average Reliability"]

    left_next = write_dataframe(
        ws,
        source_df,
        start_row=5,
        start_col=1,
        table_name=None,
    )

    ws.cell(row=4, column=5).value = "🏛 Patent Office Summary"
    ws.cell(row=4, column=5).font = Font(bold=True, color=COLORS["dark_blue"], size=12)

    patent_office_df = patent_office_summary.copy()

    if not patent_office_df.empty:
        patent_office_df.columns = ["Patent Office", "Total Patents", "Average Reliability"]
        right_next = write_dataframe(
            ws,
            patent_office_df,
            start_row=5,
            start_col=5,
            table_name=None,
        )
    else:
        right_next = 8

    section_row = max(left_next, right_next) + 2

    ws.cell(row=section_row, column=1).value = "🔬 Literature Analysis"
    ws.cell(row=section_row, column=1).font = Font(
        bold=True, color=COLORS["dark_blue"], size=12
    )

    literature_df = literature_summary.copy()

    if not literature_df.empty:
        literature_df.columns = [
            "Source Type",
            "Category",
            "Total Literature",
            "Average Reliability",
        ]
        left_next = write_dataframe(
            ws,
            literature_df.head(25),
            start_row=section_row + 1,
            start_col=1,
            table_name=None,
        )

    ws.cell(row=section_row, column=6).value = "⚙️ Maturity Analysis"
    ws.cell(row=section_row, column=6).font = Font(
        bold=True, color=COLORS["dark_blue"], size=12
    )

    maturity_df = maturity_summary.copy()

    if not maturity_df.empty:
        maturity_df.columns = [
            "Technology Maturity Level",
            "Total Sources",
            "Average Reliability",
        ]
        right_next = write_dataframe(
            ws,
            maturity_df,
            start_row=section_row + 1,
            start_col=6,
            table_name=None,
        )

    patent_row = max(left_next, right_next) + 3

    ws.cell(row=patent_row, column=1).value = "🏆 Top Patent Reference Technologies"
    ws.cell(row=patent_row, column=1).font = Font(
        bold=True, color=COLORS["dark_blue"], size=12
    )

    patent_df = patent_summary.copy().head(35)

    if not patent_df.empty:
        patent_df.columns = [
            "Technology ID",
            "Technology Name",
            "Category",
            "Patent Reference Count",
            "Average Patent Reliability",
        ]

        write_dataframe(
            ws,
            patent_df,
            start_row=patent_row + 1,
            start_col=1,
            table_name=None,
        )

    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 28
    ws.column_dimensions["G"].width = 22
    ws.column_dimensions["H"].width = 22


def create_statistical_summary(wb, statistical_summary_df):
    ws = wb.create_sheet("📈 Statistical Summary")

    style_title(
        ws,
        "Statistical Summary & Key Metrics",
        "Overall Dataset KPIs and Scoring Distributions",
        start_col=1,
        end_col=4,
    )

    stats = statistical_summary_df.iloc[0].to_dict()

    rows = [
        ("Total Technologies Tracked", stats.get("total_technologies", 0)),
        ("Average Priority Score", stats.get("average_priority_score", 0)),
        ("Average Risk Score", stats.get("average_risk_score", 0)),
        ("Average Market Readiness Score", stats.get("average_market_readiness_score", 0)),
        ("Average Adoption Potential", stats.get("average_adoption_potential_score", 0)),
        ("Critical Priority Technologies", stats.get("critical_technologies", 0)),
        ("High Priority Technologies", stats.get("high_priority_technologies", 0)),
        ("Medium Priority Technologies", stats.get("medium_priority_technologies", 0)),
        ("Low Priority Technologies", stats.get("low_priority_technologies", 0)),
        ("Highest Priority Technology", stats.get("highest_priority_technology", "")),
        ("Adoption-Priority Correlation", stats.get("adoption_priority_correlation", 0)),
    ]

    df = pd.DataFrame(rows, columns=["Metric", "Value"])
    write_dataframe(ws, df, start_row=3, start_col=1, table_name=None)

    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 42


def create_visual_charts_sheet(wb):
    ws = wb.create_sheet("📊 Visual Charts")

    style_title(
        ws,
        "Python Visualization Gallery",
        "Generated Charts for Market Trends, Risk, Priority, Competitor Activity, and Source Intelligence",
        start_col=1,
        end_col=10,
    )

    chart_files = [
        ("charts/top_technologies.png", "Top Technologies by Priority Score"),
        ("charts/category_trends.png", "Category-wise Technology Trends"),
        ("charts/competitor_activity.png", "Competitor Activity Comparison"),
        ("charts/risk_distribution.png", "Risk Distribution"),
        ("charts/research_source_distribution.png", "Research Source Distribution"),
        ("charts/patent_reference_count.png", "Patent Reference Count"),
        ("charts/release_timeline.png", "Release Timeline"),
        ("charts/risk_vs_priority.png", "Risk vs Priority Scatter Plot"),
        ("charts/patent_office_distribution.png", "Patent Office Distribution"),
        ("charts/literature_source_distribution.png", "Literature Source Distribution"),
    ]

    positions = [
        ("A4", "A3"),
        ("J4", "J3"),
        ("A27", "A26"),
        ("J27", "J26"),
        ("A50", "A49"),
        ("J50", "J49"),
        ("A73", "A72"),
        ("J73", "J72"),
        ("A96", "A95"),
        ("J96", "J95"),
    ]

    for (chart_path, chart_title), (image_cell, title_cell) in zip(chart_files, positions):
        ws[title_cell] = chart_title
        ws[title_cell].font = Font(bold=True, color=COLORS["dark_blue"], size=12)

        if os.path.exists(chart_path):
            img = Image(chart_path)
            img.width = 560
            img.height = 300
            ws.add_image(img, image_cell)
        else:
            ws[image_cell] = f"Chart file not found: {chart_path}"

    ws.sheet_view.showGridLines = False

    for col in range(1, 20):
        ws.column_dimensions[get_column_letter(col)].width = 13


def generate_excel_report():
    ensure_reports_dir()

    (
        analysis_df,
        technologies_df,
        competitors_df,
        activity_df,
        sources_df,
    ) = build_priority_analysis()

    competitor_detail_df, competitor_summary_df = build_competitor_analysis(
        activity_df, competitors_df, technologies_df
    )

    statistical_summary_df = build_statistical_summary(analysis_df)

    (
        source_type_summary,
        patent_summary,
        maturity_summary,
        patent_office_summary,
        literature_summary,
        patent_year_trend,
        category_patent_coverage,
        technology_source_summary,
    ) = build_research_source_analysis(sources_df, technologies_df)

    priority_report = analysis_df[
        [
            "technology_id",
            "technology_name",
            "category",
            "expected_release_date",
            "product_impact_score",
            "adoption_potential_score",
            "market_readiness_score",
            "rd_investment_score",
            "risk_score",
            "competitor_activity_score",
            "research_reliability_score",
            "patent_reference_count",
            "final_priority_score",
            "priority_level",
            "recommendation",
        ]
    ].sort_values("final_priority_score", ascending=False)

    risk_report = analysis_df[
        [
            "technology_name",
            "category",
            "risk_score",
            "final_priority_score",
            "priority_level",
            "recommendation",
        ]
    ].sort_values("risk_score", ascending=False)

    category_summary = (
        analysis_df.groupby("category")
        .agg(
            total_technologies=("technology_id", "count"),
            average_priority_score=("final_priority_score", "mean"),
            average_risk_score=("risk_score", "mean"),
            average_product_impact=("product_impact_score", "mean"),
            average_adoption_potential=("adoption_potential_score", "mean"),
        )
        .reset_index()
        .sort_values("average_priority_score", ascending=False)
    )

    for col in [
        "average_priority_score",
        "average_risk_score",
        "average_product_impact",
        "average_adoption_potential",
    ]:
        category_summary[col] = category_summary[col].round(2)

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    create_executive_dashboard(
        wb,
        analysis_df,
        category_summary,
        priority_report.head(10),
        statistical_summary_df,
    )

    create_priority_report(wb, priority_report)
    create_competitor_analysis(wb, competitor_summary_df)
    create_risk_analysis(wb, risk_report)
    create_patent_source_sheet(
        wb,
        source_type_summary,
        patent_summary,
        patent_office_summary,
        literature_summary,
        maturity_summary,
    )
    create_statistical_summary(wb, statistical_summary_df)
    create_visual_charts_sheet(wb)

    wb.save(OUTPUT_FILE)

    print(f"Professional Excel report generated successfully: {OUTPUT_FILE}")

    return OUTPUT_FILE